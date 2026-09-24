"""
PPTX de tendencia: evolución de Stock (Q y UF) y Asignaciones semanales por
área y total de la gerencia, a partir del histórico armado en
tablero_historico.py.

Reutiliza el logo y la paleta de colores (navy + teal) de assets/logo_jpv.png
y ficha_caso.py, para que quede consistente con el resto de los reportes
del ecosistema JPV (Reporte de Cartera, Ficha de Caso).
"""
import io
import os
from xml.sax.saxutils import escape as _escapar_xml

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Emu, Inches, Pt

LOGO_PATH = os.path.join(os.path.dirname(__file__), "assets", "logo_jpv.png")

NAVY_OSCURO = RGBColor(0x0D, 0x1F, 0x38)
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
TEAL = RGBColor(0x14, 0xA8, 0xA0)
TEAL_OSCURO = RGBColor(0x0D, 0x73, 0x77)
GRIS_TEXTO = RGBColor(0x33, 0x33, 0x33)
GRIS_CLARO = RGBColor(0xF4, 0xF6, 0xF8)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
VERDE = RGBColor(0x2E, 0xA0, 0x4E)
ROJO = RGBColor(0xC0, 0x39, 0x2B)

_NS_C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
_NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _qn(tag):
    pfx, local = tag.split(":")
    ns = {"c": _NS_C, "a": _NS_A}[pfx]
    return f"{{{ns}}}{local}"

ANCHO_SLIDE = Inches(13.333)
ALTO_SLIDE = Inches(7.5)
HEADER_ALTO = Inches(1.1)


def _agregar_header(slide, titulo, subtitulo):
    barra = slide.shapes.add_shape(1, 0, 0, ANCHO_SLIDE, HEADER_ALTO)
    barra.fill.solid()
    barra.fill.fore_color.rgb = NAVY
    barra.line.fill.background()
    barra.shadow.inherit = False

    divisor = slide.shapes.add_shape(1, 0, HEADER_ALTO, ANCHO_SLIDE, Emu(45720))
    divisor.fill.solid()
    divisor.fill.fore_color.rgb = TEAL_OSCURO
    divisor.line.fill.background()
    divisor.shadow.inherit = False

    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(LOGO_PATH, Inches(0.35), Inches(0.2), height=Inches(0.7))

    caja_titulo = slide.shapes.add_textbox(Inches(2.2), Inches(0.18), ANCHO_SLIDE - Inches(2.6), Inches(0.85))
    tf = caja_titulo.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = titulo
    p.font.size, p.font.bold, p.font.color.rgb = Pt(26), True, BLANCO
    if subtitulo:
        p2 = tf.add_paragraph()
        p2.text = subtitulo
        p2.font.size, p2.font.italic, p2.font.color.rgb = Pt(13), True, TEAL


def _slide_base(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fondo = slide.background.fill
    fondo.solid()
    fondo.fore_color.rgb = BLANCO
    return slide


def _agregar_titulo_chico(slide, texto, left, top, width):
    caja = slide.shapes.add_textbox(left, top, width, Inches(0.35))
    p = caja.text_frame.paragraphs[0]
    p.text = texto
    p.font.size, p.font.bold, p.font.color.rgb = Pt(14), True, NAVY


def _agregar_grafico(slide, tipo, categorias, series, num_format=None,
                      left=Inches(0.5), top=None, width=None, height=None,
                      leyenda=True, color_offset=0):
    if top is None:
        top = HEADER_ALTO + Inches(0.35)
    if width is None:
        width = ANCHO_SLIDE - Inches(1)
    if height is None:
        height = ALTO_SLIDE - HEADER_ALTO - Inches(0.85)

    cd = CategoryChartData()
    cd.categories = categorias
    for nombre, valores in series:
        cd.add_series(nombre, valores)

    grafico_frame = slide.shapes.add_chart(tipo, left, top, width, height, cd)
    chart = grafico_frame.chart
    chart.has_legend = leyenda
    if leyenda:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(12)
        chart.legend.font.color.rgb = GRIS_TEXTO

    try:
        eje_cat = chart.category_axis
        eje_cat.tick_labels.font.size = Pt(8)
        eje_cat.tick_labels.font.color.rgb = GRIS_TEXTO
    except Exception:
        pass
    try:
        eje_valor = chart.value_axis
        eje_valor.tick_labels.font.size = Pt(10)
        eje_valor.tick_labels.font.color.rgb = GRIS_TEXTO
        if num_format:
            eje_valor.tick_labels.number_format = num_format
            eje_valor.tick_labels.number_format_is_linked = False
    except Exception:
        pass

    colores = [TEAL, NAVY, TEAL_OSCURO]
    for i, serie in enumerate(chart.series):
        color = colores[(i + color_offset) % len(colores)]
        try:
            if tipo in (XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_MARKERS):
                serie.format.line.color.rgb = color
                serie.format.line.width = Pt(2.5)
                serie.marker.format.fill.solid()
                serie.marker.format.fill.fore_color.rgb = color
            else:
                serie.format.fill.solid()
                serie.format.fill.fore_color.rgb = color
        except Exception:
            pass
    return chart


def _agregar_par_dividido(slide, tipo, etiquetas, serie_ing, serie_movil, num_format=None):
    """Dos gráficos lado a lado (Ingeniería / Equipo Móvil), cada uno con su
    propio eje: cuando la escala de una división es varias veces la de la
    otra, un solo eje combinado aplana la tendencia de la más chica."""
    mitad = Emu((ANCHO_SLIDE - Inches(1.2)) // 2)
    top_grafico = HEADER_ALTO + Inches(0.75)
    alto_grafico = ALTO_SLIDE - top_grafico - Inches(0.4)

    _agregar_titulo_chico(slide, "Ingeniería y Energía", Inches(0.5), HEADER_ALTO + Inches(0.3), mitad)
    _agregar_grafico(slide, tipo, etiquetas, [("Ingeniería y Energía", serie_ing)],
                      num_format=num_format, left=Inches(0.5), top=top_grafico, width=mitad, height=alto_grafico,
                      leyenda=False, color_offset=0)

    _agregar_titulo_chico(slide, "Equipo Móvil", Inches(0.7) + mitad, HEADER_ALTO + Inches(0.3), mitad)
    _agregar_grafico(slide, tipo, etiquetas, [("Equipo Móvil", serie_movil)],
                      num_format=num_format, left=Inches(0.7) + mitad, top=top_grafico, width=mitad, height=alto_grafico,
                      leyenda=False, color_offset=1)


def _agregar_grafico_combo(slide, etiquetas, valores_barras, nombre_barras, valores_linea, nombre_linea,
                            left=None, top=None, width=None, height=None, num_format_linea=None):
    """Barras (coloreadas verde/rojo según el signo) + una línea en un eje
    SECUNDARIO — para mostrar, en un mismo gráfico, el Neto semanal
    (Ingreso − Egreso) junto con el nivel de Stock que ese neto explica.
    python-pptx no arma gráficos combinados de forma nativa: se inyecta el
    lineChart y el eje secundario directo en el XML del gráfico (es el mismo
    patrón que usa Excel/PowerPoint para un "combo chart")."""
    if left is None:
        left = Inches(0.5)
    if top is None:
        top = HEADER_ALTO + Inches(0.35)
    if width is None:
        width = ANCHO_SLIDE - Inches(1)
    if height is None:
        height = ALTO_SLIDE - HEADER_ALTO - Inches(0.85)

    cd = CategoryChartData()
    cd.categories = etiquetas
    cd.add_series(nombre_barras, valores_barras)
    grafico_frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, left, top, width, height, cd)
    chart = grafico_frame.chart

    serie_barras = chart.plots[0].series[0]
    for punto, valor in zip(serie_barras.points, valores_barras):
        punto.format.fill.solid()
        punto.format.fill.fore_color.rgb = ROJO if (valor or 0) < 0 else VERDE

    chart_space = chart._chartSpace
    plot_area = chart_space.find(_qn("c:chart")).find(_qn("c:plotArea"))
    bar_chart_el = plot_area.find(_qn("c:barChart"))
    cat_ax_el = plot_area.find(_qn("c:catAx"))
    val_ax_el = plot_area.find(_qn("c:valAx"))

    cat_ax_id_1 = cat_ax_el.find(_qn("c:axId")).get("val")
    val_ax_id_1 = val_ax_el.find(_qn("c:axId")).get("val")
    cat_ax_id_2 = str(int(cat_ax_id_1) + 1)
    val_ax_id_2 = str(int(val_ax_id_1) + 1)

    formato = num_format_linea or "General"
    color_linea = str(TEAL_OSCURO)
    pts_cat = "".join(f'<c:pt idx="{i}"><c:v>{_escapar_xml(str(c))}</c:v></c:pt>' for i, c in enumerate(etiquetas))
    pts_val = "".join(
        f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in enumerate(valores_linea) if v is not None
    )

    line_chart_xml = f'''<c:lineChart xmlns:c="{_NS_C}" xmlns:a="{_NS_A}">
  <c:grouping val="standard"/>
  <c:varyColors val="0"/>
  <c:ser>
    <c:idx val="1"/>
    <c:order val="1"/>
    <c:tx><c:strRef><c:f>Sheet1!$C$1</c:f><c:strCache><c:ptCount val="1"/><c:pt idx="0"><c:v>{_escapar_xml(nombre_linea)}</c:v></c:pt></c:strCache></c:strRef></c:tx>
    <c:spPr><a:ln w="28575"><a:solidFill><a:srgbClr val="{color_linea}"/></a:solidFill></a:ln></c:spPr>
    <c:marker><c:symbol val="circle"/><c:size val="6"/><c:spPr><a:solidFill><a:srgbClr val="{color_linea}"/></a:solidFill></c:spPr></c:marker>
    <c:cat><c:strRef><c:f>Sheet1!$A$2:$A${1 + len(etiquetas)}</c:f><c:strCache><c:ptCount val="{len(etiquetas)}"/>{pts_cat}</c:strCache></c:strRef></c:cat>
    <c:val><c:numRef><c:f>Sheet1!$C$2:$C${1 + len(etiquetas)}</c:f><c:numCache><c:formatCode>{formato}</c:formatCode><c:ptCount val="{len(etiquetas)}"/>{pts_val}</c:numCache></c:numRef></c:val>
    <c:smooth val="0"/>
  </c:ser>
  <c:marker val="1"/>
  <c:axId val="{cat_ax_id_2}"/>
  <c:axId val="{val_ax_id_2}"/>
</c:lineChart>'''
    bar_chart_el.addnext(etree.fromstring(line_chart_xml))

    val_ax2_xml = f'''<c:valAx xmlns:c="{_NS_C}" xmlns:a="{_NS_A}">
  <c:axId val="{val_ax_id_2}"/>
  <c:scaling><c:orientation val="minMax"/></c:scaling>
  <c:delete val="0"/>
  <c:axPos val="r"/>
  <c:numFmt formatCode="{formato}" sourceLinked="0"/>
  <c:majorTickMark val="out"/>
  <c:minorTickMark val="none"/>
  <c:tickLblPos val="nextTo"/>
  <c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="1000"><a:solidFill><a:srgbClr val="333333"/></a:solidFill></a:defRPr></a:pPr><a:endParaRPr lang="es-CL"/></a:p></c:txPr>
  <c:crossAx val="{cat_ax_id_2}"/>
  <c:crosses val="max"/>
</c:valAx>'''
    val_ax2 = etree.fromstring(val_ax2_xml)

    cat_ax2_xml = f'''<c:catAx xmlns:c="{_NS_C}">
  <c:axId val="{cat_ax_id_2}"/>
  <c:scaling><c:orientation val="minMax"/></c:scaling>
  <c:delete val="1"/>
  <c:axPos val="b"/>
  <c:majorTickMark val="out"/>
  <c:minorTickMark val="none"/>
  <c:tickLblPos val="nextTo"/>
  <c:crossAx val="{val_ax_id_2}"/>
  <c:crosses val="autoZero"/>
  <c:auto val="1"/>
  <c:lblAlgn val="ctr"/>
  <c:lblOffset val="100"/>
  <c:noMultiLvlLbl val="0"/>
</c:catAx>'''
    cat_ax2 = etree.fromstring(cat_ax2_xml)

    val_ax_el.addnext(cat_ax2)
    cat_ax2.addnext(val_ax2)

    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(12)
    chart.legend.font.color.rgb = GRIS_TEXTO

    try:
        chart.category_axis.tick_labels.font.size = Pt(8)
        chart.category_axis.tick_labels.font.color.rgb = GRIS_TEXTO
    except Exception:
        pass
    try:
        chart.value_axis.tick_labels.font.size = Pt(10)
        chart.value_axis.tick_labels.font.color.rgb = GRIS_TEXTO
    except Exception:
        pass

    return chart


def _agregar_par_dividido_combo(slide, etiquetas, datos_ing, datos_mov, nombre_barras, nombre_linea, num_format_linea=None):
    """Dos gráficos combo (barras + línea en eje secundario) lado a lado,
    uno por división. datos_ing / datos_mov: (valores_barras, valores_linea)."""
    mitad = Emu((ANCHO_SLIDE - Inches(1.2)) // 2)
    top_grafico = HEADER_ALTO + Inches(0.75)
    alto_grafico = ALTO_SLIDE - top_grafico - Inches(0.4)

    _agregar_titulo_chico(slide, "Ingeniería y Energía", Inches(0.5), HEADER_ALTO + Inches(0.3), mitad)
    _agregar_grafico_combo(slide, etiquetas, datos_ing[0], nombre_barras, datos_ing[1], nombre_linea,
                            left=Inches(0.5), top=top_grafico, width=mitad, height=alto_grafico,
                            num_format_linea=num_format_linea)

    _agregar_titulo_chico(slide, "Equipo Móvil", Inches(0.7) + mitad, HEADER_ALTO + Inches(0.3), mitad)
    _agregar_grafico_combo(slide, etiquetas, datos_mov[0], nombre_barras, datos_mov[1], nombre_linea,
                            left=Inches(0.7) + mitad, top=top_grafico, width=mitad, height=alto_grafico,
                            num_format_linea=num_format_linea)


def _serie_neto(serie_a, serie_b):
    return [None if (a is None or b is None) else a - b for a, b in zip(serie_a, serie_b)]


MESES_ES = {1: "ene", 2: "feb", 3: "mar", 4: "abr", 5: "may", 6: "jun",
            7: "jul", 8: "ago", 9: "sep", 10: "oct", 11: "nov", 12: "dic"}


def _promedio_semanal_y_percapita(df_historico, df_promedio, columna, fechas):
    """(promedio semanal total, promedio semanal por ajustador vigente) de
    `columna` para 'Total Gerencia' en el rango `fechas`. El per-cápita se
    calcula semana a semana (valor de la semana / ajustadores vigentes esa
    semana, excluyendo a Dagoberto Pérez) y luego se promedia — no divide el
    promedio total por el promedio de ajustadores."""
    sub_valor = df_historico[df_historico["Division"] == "Total Gerencia"].set_index("Fecha")[columna]
    valores = [sub_valor.get(f) for f in fechas]
    valores_validos = [v for v in valores if v is not None]
    promedio_total = sum(valores_validos) / len(valores_validos) if valores_validos else None

    promedio_percapita = None
    if df_promedio is not None and not df_promedio.empty:
        sub_n = df_promedio[df_promedio["Division"] == "Total Gerencia"].set_index("Fecha")["N_Ajustadores"]
        razones = [v / sub_n[f] for f, v in zip(fechas, valores) if v is not None and sub_n.get(f)]
        promedio_percapita = sum(razones) / len(razones) if razones else None
    return promedio_total, promedio_percapita


def _subtitulo_promedio(df_historico, df_promedio, columna, fechas, unidad):
    """'Promedio semanal: X <unidad> · Y <unidad> por ajustador vigente',
    o None si no hay datos — para poner en el subtítulo del encabezado."""
    promedio_total, promedio_percapita = _promedio_semanal_y_percapita(df_historico, df_promedio, columna, fechas)
    partes = []
    if promedio_total is not None:
        partes.append(f"Promedio semanal: {promedio_total:,.0f} {unidad}")
    if promedio_percapita is not None:
        partes.append(f"{promedio_percapita:,.0f} {unidad} por ajustador vigente")
    return " · ".join(partes) or None


def generar_pptx_historico(df_historico, df_promedio=None):
    """df_historico: salida de tablero_historico.parsear_historico().
    df_promedio: salida opcional de
    tablero_historico.parsear_promedio_casos_por_ajustador() — si se pasa,
    agrega una slide con la evolución del Stock (Q) promedio por ajustador."""
    if df_historico is None or df_historico.empty:
        raise ValueError("No hay datos históricos para generar el PPTX.")

    prs = Presentation()
    prs.slide_width = ANCHO_SLIDE
    prs.slide_height = ALTO_SLIDE

    fechas = sorted(df_historico["Fecha"].unique())
    etiquetas = [f"{f.day:02d}-{MESES_ES[f.month]}" for f in fechas]

    def serie_de(division, columna):
        sub = df_historico[df_historico["Division"] == division].set_index("Fecha")[columna]
        return [sub.get(f, None) for f in fechas]

    rango = f"{etiquetas[0]} al {etiquetas[-1]} de {fechas[-1].year} · {len(fechas)} semanas"

    slide = _slide_base(prs)
    _agregar_header(slide, "Tablero Gerencial — Evolución Histórica", f"Ingeniería y Equipo Móvil · {rango}")
    caja = slide.shapes.add_textbox(Inches(0.5), Inches(3), ANCHO_SLIDE - Inches(1), Inches(2))
    tf = caja.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "JPV Asociados — Ajustadores Especializados"
    p.font.size, p.font.bold, p.font.color.rgb = Pt(18), True, NAVY

    slide = _slide_base(prs)
    _agregar_header(slide, "Evolución de Stock — Honorarios (UF)", "Total Gerencia")
    _agregar_grafico(slide, XL_CHART_TYPE.LINE_MARKERS, etiquetas, [
        ("Total Gerencia", serie_de("Total Gerencia", "Stock_UF")),
    ], num_format="#,##0", color_offset=2)

    slide = _slide_base(prs)
    _agregar_header(slide, "Evolución de Stock — Cantidad de Casos (Q)", "Total Gerencia")
    _agregar_grafico(slide, XL_CHART_TYPE.LINE_MARKERS, etiquetas, [
        ("Total Gerencia", serie_de("Total Gerencia", "Stock_Q")),
    ], color_offset=2)

    slide = _slide_base(prs)
    _agregar_header(slide, "Evolución de Stock — Cantidad de Casos (Q)", None)
    _agregar_grafico(slide, XL_CHART_TYPE.LINE_MARKERS, etiquetas, [
        ("Ingeniería y Energía", serie_de("Ingeniería y Energía", "Stock_Q")),
        ("Equipo Móvil", serie_de("Equipo Móvil", "Stock_Q")),
    ])

    slide = _slide_base(prs)
    _agregar_header(slide, "Evolución de Stock — Honorarios (UF)", None)
    # Un gráfico por división: la escala de Ingeniería en UF es muchas veces
    # la de Equipo Móvil, y un solo eje combinado aplanaba su tendencia.
    _agregar_par_dividido(slide, XL_CHART_TYPE.LINE_MARKERS, etiquetas,
                           serie_de("Ingeniería y Energía", "Stock_UF"), serie_de("Equipo Móvil", "Stock_UF"),
                           num_format="#,##0")

    slide = _slide_base(prs)
    _agregar_header(slide, "Asignaciones Semanales por Área",
                     _subtitulo_promedio(df_historico, df_promedio, "Asignados", fechas, "casos"))
    _agregar_par_dividido(slide, XL_CHART_TYPE.COLUMN_CLUSTERED, etiquetas,
                           serie_de("Ingeniería y Energía", "Asignados"), serie_de("Equipo Móvil", "Asignados"))

    slide = _slide_base(prs)
    _agregar_header(slide, "Asignaciones Semanales — Total Gerencia",
                     _subtitulo_promedio(df_historico, df_promedio, "Asignados", fechas, "casos"))
    _agregar_grafico(slide, XL_CHART_TYPE.COLUMN_CLUSTERED, etiquetas, [
        ("Total Gerencia", serie_de("Total Gerencia", "Asignados")),
    ], leyenda=False)

    slide = _slide_base(prs)
    _agregar_header(slide, "IFL Emitidos — Cantidad de Casos",
                     _subtitulo_promedio(df_historico, df_promedio, "IFL_Q", fechas, "casos"))
    _agregar_par_dividido(slide, XL_CHART_TYPE.COLUMN_CLUSTERED, etiquetas,
                           serie_de("Ingeniería y Energía", "IFL_Q"), serie_de("Equipo Móvil", "IFL_Q"))

    slide = _slide_base(prs)
    _agregar_header(slide, "IFL Emitidos — Total Gerencia",
                     _subtitulo_promedio(df_historico, df_promedio, "IFL_Q", fechas, "casos"))
    _agregar_grafico(slide, XL_CHART_TYPE.LINE_MARKERS, etiquetas, [
        ("Total Gerencia", serie_de("Total Gerencia", "IFL_Q")),
    ], leyenda=False, color_offset=2)

    slide = _slide_base(prs)
    _agregar_header(slide, "IFL Emitidos — Honorarios (UF)",
                     _subtitulo_promedio(df_historico, df_promedio, "IFL_UF", fechas, "UF"))
    # Un gráfico por división: la escala de Ingeniería en UF es muchas veces
    # la de Equipo Móvil, igual que en Stock UF.
    _agregar_par_dividido(slide, XL_CHART_TYPE.COLUMN_CLUSTERED, etiquetas,
                           serie_de("Ingeniería y Energía", "IFL_UF"), serie_de("Equipo Móvil", "IFL_UF"),
                           num_format="#,##0")

    slide = _slide_base(prs)
    _agregar_header(slide, "IFL Emitidos — Total Gerencia",
                     _subtitulo_promedio(df_historico, df_promedio, "IFL_UF", fechas, "UF"))
    _agregar_grafico(slide, XL_CHART_TYPE.LINE_MARKERS, etiquetas, [
        ("Total Gerencia", serie_de("Total Gerencia", "IFL_UF")),
    ], num_format="#,##0", leyenda=False, color_offset=2)

    slide = _slide_base(prs)
    _agregar_header(slide, "Ingresos vs Egresos — Total Gerencia", "Barras: Neto semanal (Asignaciones − IFL) · Línea: Stock (Q)")
    neto_gerencia = _serie_neto(serie_de("Total Gerencia", "Asignados"), serie_de("Total Gerencia", "IFL_Q"))
    _agregar_grafico_combo(slide, etiquetas, neto_gerencia, "Neto (Ingreso − Egreso)",
                            serie_de("Total Gerencia", "Stock_Q"), "Stock (Q)")

    slide = _slide_base(prs)
    _agregar_header(slide, "Ingresos vs Egresos por Área", "Barras: Neto semanal (Asignaciones − IFL) · Línea: Stock (Q)")
    neto_ing = _serie_neto(serie_de("Ingeniería y Energía", "Asignados"), serie_de("Ingeniería y Energía", "IFL_Q"))
    neto_mov = _serie_neto(serie_de("Equipo Móvil", "Asignados"), serie_de("Equipo Móvil", "IFL_Q"))
    _agregar_par_dividido_combo(
        slide, etiquetas,
        (neto_ing, serie_de("Ingeniería y Energía", "Stock_Q")),
        (neto_mov, serie_de("Equipo Móvil", "Stock_Q")),
        "Neto (Ingreso − Egreso)", "Stock (Q)",
    )

    if df_promedio is not None and not df_promedio.empty:
        fechas_prom = sorted(df_promedio["Fecha"].unique())
        etiquetas_prom = [f"{f.day:02d}-{MESES_ES[f.month]}" for f in fechas_prom]

        def serie_promedio(division):
            sub = df_promedio[df_promedio["Division"] == division].set_index("Fecha")["Promedio_Q"]
            return [sub.get(f, None) for f in fechas_prom]

        slide = _slide_base(prs)
        _agregar_header(slide, "Stock Promedio por Ajustador", "Excluye a Dagoberto Pérez (carga no representativa)")
        _agregar_grafico(slide, XL_CHART_TYPE.LINE_MARKERS, etiquetas_prom, [
            ("Ingeniería y Energía", serie_promedio("Ingeniería y Energía")),
            ("Equipo Móvil", serie_promedio("Equipo Móvil")),
        ])

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
